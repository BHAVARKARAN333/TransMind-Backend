import os
import tempfile
import re
from docx import Document

debug_path = os.path.join(tempfile.gettempdir(), "transmind_last_upload.docx")

if os.path.exists(debug_path):
    doc = Document(debug_path)
    for para in doc.paragraphs:
        t = para.text.strip()
        if "Highly motivated" in t:
            orig1 = "Highly motivated Bachelor of Computer Science student (2025) with strong foundation in Java, and Database Management Systems."
            orig2 = "Experienced in developing academic projects including a Tuition Management System."
            
            w1 = orig1.split()
            p1 = r'[\s\u200b\u00A0]*'.join(re.escape(w) for w in w1)
            text1, c1 = re.subn(p1, "HINDI_1", para.text, count=1)
            
            w2 = orig2.split()
            p2 = r'[\s\u200b\u00A0]*'.join(re.escape(w) for w in w2)
            final_text, c2 = re.subn(p2, "HINDI_2", text1, count=1)
            
            print("Final combined text string:", repr(final_text))
            
            # Now let's try calling docx_exporter's exact code
            import sys
            sys.path.append(os.path.dirname(__file__))
            from docx_exporter import _replace_para_text
            
            _replace_para_text(para, final_text)
            
    # Save test output
    outpath = "regex_test_output.docx"
    doc.save(outpath)
    
    # Reload and see what .text actually became
    doc2 = Document(outpath)
    for para in doc2.paragraphs:
        if "HINDI_1" in para.text or "Highly motivated" in para.text:
            print("RELOADED TEXT:", repr(para.text))
            break
